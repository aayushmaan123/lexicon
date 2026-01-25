"""DOCX document parser using python-docx."""

import os
from typing import Any

from docx import Document as DocxDocument

from lexicon.domain.entities.document import DocumentType, ParsedDocument, Section
from lexicon.domain.parsers.base import DocumentParser


class DOCXParser(DocumentParser):
    """DOCX document parser using python-docx."""

    def parse(self, file_path: str) -> ParsedDocument:
        """Parse a DOCX document and return structured data.

        Args:
            file_path: Path to the DOCX file

        Returns:
            ParsedDocument with extracted text and structure

        Raises:
            FileNotFoundError: If file does not exist
            ValueError: If file is not a valid DOCX
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        try:
            doc = DocxDocument(file_path)
        except Exception as e:
            raise ValueError(f"Invalid DOCX file: {e}") from e

        try:
            full_text = []
            sections = []
            current_section = None
            current_section_text = []

            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue

                # Check if paragraph is a heading
                if para.style.name.startswith("Heading"):
                    # Save previous section
                    if current_section and current_section_text:
                        current_section.content = "\n".join(current_section_text)
                        sections.append(current_section)

                    # Extract heading level (e.g., "Heading 1" -> 1)
                    try:
                        level = int(para.style.name.split()[-1])
                    except (ValueError, IndexError):
                        level = 1

                    # Start new section
                    current_section = Section(
                        title=text,
                        content="",
                        level=level,
                    )
                    current_section_text = []
                    full_text.append(text)
                else:
                    # Regular paragraph
                    full_text.append(text)
                    if current_section:
                        current_section_text.append(text)

            # Save last section
            if current_section and current_section_text:
                current_section.content = "\n".join(current_section_text)
                sections.append(current_section)

            metadata = self.extract_metadata(file_path)

            return ParsedDocument(
                document_id=os.path.basename(file_path),
                filename=os.path.basename(file_path),
                document_type=DocumentType.GENERAL,
                full_text="\n\n".join(full_text),
                sections=sections,
                metadata=metadata,
            )

        except Exception as e:
            raise ValueError(f"Error parsing DOCX: {e}") from e

    def extract_metadata(self, file_path: str) -> dict[str, Any]:
        """Extract metadata from a DOCX document.

        Args:
            file_path: Path to the DOCX file

        Returns:
            Dictionary containing DOCX metadata

        Raises:
            FileNotFoundError: If file does not exist
            ValueError: If file is not a valid DOCX
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        try:
            doc = DocxDocument(file_path)
            core_props = doc.core_properties

            return {
                "author": core_props.author or "",
                "title": core_props.title or "",
                "subject": core_props.subject or "",
                "created": str(core_props.created) if core_props.created else "",
                "modified": str(core_props.modified) if core_props.modified else "",
                "last_modified_by": core_props.last_modified_by or "",
                "category": core_props.category or "",
                "comments": core_props.comments or "",
            }

        except Exception as e:
            raise ValueError(f"Error extracting DOCX metadata: {e}") from e
