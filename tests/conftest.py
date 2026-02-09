"""Pytest configuration and shared fixtures for Lexicon tests."""

import os
import tempfile
from datetime import datetime
from pathlib import Path
from typing import AsyncGenerator, Generator
from unittest.mock import AsyncMock, Mock, patch

import pytest
from pydantic import BaseModel

from lexicon.domain.entities.contract import Contract, Clause, RiskLevel, Party, PartyType
from lexicon.domain.entities.document import Document, DocumentType, ParsedDocument, Section
from lexicon.domain.entities.legal_research import CaseLaw, ResearchQuery, Statute
from lexicon.shared.config import Settings


# ============================================================================
# Configuration Fixtures
# ============================================================================


@pytest.fixture
def test_settings() -> Settings:
    """Provide test configuration settings."""
    return Settings(
        openai_api_key="test-openai-key",
        anthropic_api_key="test-anthropic-key",
        database_url="sqlite:///:memory:",
        redis_url="redis://localhost:6379/1",
        chroma_persist_directory="./test_chroma_data",
        chroma_collection_name="test_collection",
        primary_llm_model="gpt-4o",
        secondary_llm_model="gpt-4o-mini",
        fallback_llm_model="claude-3-5-sonnet-20241022",
        embedding_model="text-embedding-3-large",
        log_level="DEBUG",
        environment="test",
        max_retries=2,
        retry_backoff_factor=2,
        chunk_size_contract=512,
        chunk_overlap_contract=50,
    )


@pytest.fixture(autouse=True)
def mock_settings(test_settings):
    """Mock settings for all tests."""
    with patch("lexicon.shared.config.settings", test_settings):
        yield test_settings


# ============================================================================
# Mock LLM Provider Fixtures
# ============================================================================


class MockResponse(BaseModel):
    """Mock structured response for testing."""
    content: str


@pytest.fixture
def mock_openai_provider():
    """Mock OpenAI provider for testing."""
    provider = Mock()
    provider.model = "gpt-4o-mini"
    provider.generate = Mock(return_value="Mock LLM response text")
    provider.generate_structured = Mock(return_value=MockResponse(content="Mock structured response"))
    provider.count_tokens = Mock(return_value=100)
    provider.get_cost = Mock(return_value=0.001)
    return provider


@pytest.fixture
def mock_anthropic_provider():
    """Mock Anthropic provider for testing."""
    provider = Mock()
    provider.model = "claude-3-5-sonnet-20241022"
    provider.generate = Mock(return_value="Mock Claude response text")
    provider.generate_structured = Mock(return_value=MockResponse(content="Mock Claude structured response"))
    provider.count_tokens = Mock(return_value=120)
    provider.get_cost = Mock(return_value=0.002)
    return provider


@pytest.fixture
def mock_openai_client():
    """Mock OpenAI API client."""
    with patch("lexicon.infrastructure.openai_client.OpenAI") as mock_client:
        mock_completion = Mock()
        mock_completion.choices = [Mock(message=Mock(content="Mocked OpenAI response", parsed=MockResponse(content="Parsed")))]
        mock_completion.usage = Mock(prompt_tokens=50, completion_tokens=50)
        
        mock_client.return_value.chat.completions.create.return_value = mock_completion
        mock_client.return_value.beta.chat.completions.parse.return_value = mock_completion
        
        yield mock_client


@pytest.fixture
def mock_anthropic_client():
    """Mock Anthropic API client."""
    with patch("lexicon.infrastructure.anthropic_client.Anthropic") as mock_client:
        mock_message = Mock()
        mock_message.content = [Mock(text="Mocked Anthropic response")]
        mock_message.usage = Mock(input_tokens=60, output_tokens=40)
        
        mock_client.return_value.messages.create.return_value = mock_message
        
        yield mock_client


# ============================================================================
# Mock Vector Store Fixtures
# ============================================================================


@pytest.fixture
def mock_chroma_client():
    """Mock ChromaDB client."""
    with patch("lexicon.infrastructure.chroma_client.chromadb.PersistentClient") as mock_client:
        mock_collection = Mock()
        mock_collection.count.return_value = 0
        mock_collection.add = Mock()
        mock_collection.query = Mock(return_value={
            "ids": [["doc1", "doc2"]],
            "documents": [["Document 1 text", "Document 2 text"]],
            "metadatas": [[{"source": "test1"}, {"source": "test2"}]],
            "distances": [[0.1, 0.2]],
        })
        mock_collection.delete = Mock()
        
        mock_client.return_value.get_or_create_collection.return_value = mock_collection
        mock_client.return_value.create_collection.return_value = mock_collection
        mock_client.return_value.delete_collection = Mock()
        
        yield mock_client


@pytest.fixture
def mock_vector_store(mock_chroma_client):
    """Mock vector store instance."""
    from lexicon.infrastructure.chroma_client import ChromaVectorStore
    
    store = ChromaVectorStore(
        collection_name="test_collection",
        persist_directory="./test_chroma_data"
    )
    return store


# ============================================================================
# Sample Document Fixtures
# ============================================================================


@pytest.fixture
def sample_text() -> str:
    """Provide sample text for testing."""
    return """
    This is a sample legal document for testing purposes.
    
    PARTIES:
    Party A: John Doe, an individual
    Party B: Acme Corporation, a Delaware corporation
    
    AGREEMENT:
    This agreement is entered into on January 1, 2024.
    The parties agree to the following terms and conditions.
    
    Section 1: Definitions
    "Services" means the consulting services provided by Party A.
    
    Section 2: Payment
    Party B shall pay Party A the sum of $10,000 for the Services.
    """


@pytest.fixture
def sample_pdf_path(tmp_path) -> Path:
    """Create a sample PDF file for testing."""
    import fitz
    
    pdf_path = tmp_path / "sample.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Sample PDF Document\n\nThis is a test PDF file.")
    doc.save(pdf_path)
    doc.close()
    
    return pdf_path


@pytest.fixture
def sample_docx_path(tmp_path) -> Path:
    """Create a sample DOCX file for testing."""
    from docx import Document as DocxDocument
    
    docx_path = tmp_path / "sample.docx"
    doc = DocxDocument()
    doc.add_heading("Sample DOCX Document", 0)
    doc.add_paragraph("This is a test DOCX file.")
    doc.save(docx_path)
    
    return docx_path


@pytest.fixture
def sample_document() -> Document:
    """Provide a sample Document entity."""
    return Document(
        document_id="doc-123",
        filename="sample.pdf",
        file_path="/tmp/sample.pdf",
        document_type=DocumentType.CONTRACT,
        uploaded_at=datetime.now(),
        file_size=1024,
    )


@pytest.fixture
def sample_parsed_document() -> ParsedDocument:
    """Provide a sample ParsedDocument entity."""
    return ParsedDocument(
        document_id="doc-123",
        filename="sample.pdf",
        document_type=DocumentType.CONTRACT,
        full_text="This is the full text of the document.",
        sections=[
            Section(
                title="Introduction",
                content="This is the introduction section.",
                page_number=1,
                level=1,
            ),
            Section(
                title="Terms and Conditions",
                content="These are the terms and conditions.",
                page_number=2,
                level=1,
            ),
        ],
        metadata={
            "author": "Test Author",
            "title": "Sample Document",
            "page_count": 2,
        },
        page_count=2,
    )


# ============================================================================
# Contract Fixtures
# ============================================================================


@pytest.fixture
def sample_contract() -> Contract:
    """Provide a sample Contract entity."""
    return Contract(
        document_id="contract-123",
        contract_type="Service Agreement",
        parties=[
            Party(name="John Doe", type=PartyType.INDIVIDUAL, role="Consultant"),
            Party(name="Acme Corporation", type=PartyType.CORPORATION, role="Client"),
        ],
        clauses=[
            Clause(
                type="payment",
                title="Payment Terms",
                content="Payment shall be made within 30 days.",
                page_number=1,
                risk_level=RiskLevel.LOW,
            ),
            Clause(
                type="liability",
                title="Liability Limitation",
                content="Liability is limited to $100,000.",
                page_number=2,
                risk_level=RiskLevel.MEDIUM,
            ),
        ],
        effective_date=datetime(2024, 1, 1),
        expiration_date=datetime(2025, 1, 1),
        jurisdiction="Delaware",
    )


# ============================================================================
# Legal Research Fixtures
# ============================================================================


@pytest.fixture
def sample_legal_query() -> ResearchQuery:
    """Provide a sample ResearchQuery entity."""
    return ResearchQuery(
        query_text="What are the requirements for a valid contract?",
        jurisdiction="United States",
        query_type="general",
    )


@pytest.fixture
def sample_case_law() -> CaseLaw:
    """Provide a sample CaseLaw entity."""
    return CaseLaw(
        case_name="Smith v. Jones",
        citation="123 F.3d 456",
        court="9th Circuit",
        jurisdiction="Federal",
        decision_date=datetime(2020, 5, 15).date(),
        summary="Court held that contracts require mutual assent.",
        key_holdings=["Mutual assent is required", "Consideration must be present"],
    )


@pytest.fixture
def sample_statute() -> Statute:
    """Provide a sample Statute entity."""
    return Statute(
        title="Uniform Commercial Code",
        citation="U.C.C. § 2-201",
        jurisdiction="United States",
        section_number="2-201",
        summary="Statute of frauds for sale of goods.",
        effective_date=datetime(1952, 1, 1).date(),
    )


# ============================================================================
# Cache and Database Mocks
# ============================================================================


@pytest.fixture
def mock_redis_client():
    """Mock Redis client."""
    with patch("lexicon.infrastructure.cache_client.redis.Redis") as mock_redis:
        mock_instance = Mock()
        mock_instance.get = Mock(return_value=None)
        mock_instance.set = Mock(return_value=True)
        mock_instance.delete = Mock(return_value=1)
        mock_instance.exists = Mock(return_value=False)
        
        mock_redis.return_value = mock_instance
        yield mock_redis


@pytest.fixture
def mock_database():
    """Mock database connection."""
    with patch("lexicon.infrastructure.database.create_engine") as mock_engine:
        yield mock_engine


# ============================================================================
# File System Fixtures
# ============================================================================


@pytest.fixture
def temp_directory() -> Generator[Path, None, None]:
    """Provide a temporary directory for testing."""
    with tempfile.TemporaryDirectory() as tmpdir:
        yield Path(tmpdir)


@pytest.fixture
def cleanup_chroma():
    """Cleanup ChromaDB test data after tests."""
    yield
    
    # Cleanup test ChromaDB directory
    import shutil
    test_dir = Path("./test_chroma_data")
    if test_dir.exists():
        shutil.rmtree(test_dir)


# ============================================================================
# Async Fixtures
# ============================================================================


@pytest.fixture
async def async_mock_llm():
    """Mock async LLM provider."""
    provider = AsyncMock()
    provider.generate = AsyncMock(return_value="Async mock response")
    provider.generate_structured = AsyncMock(return_value=MockResponse(content="Async structured"))
    return provider


# ============================================================================
# API Testing Fixtures
# ============================================================================


@pytest.fixture
def mock_fastapi_app():
    """Mock FastAPI application for testing."""
    from fastapi.testclient import TestClient
    from lexicon.api.main import app
    
    return TestClient(app)


# ============================================================================
# CLI Testing Fixtures
# ============================================================================


@pytest.fixture
def cli_runner():
    """Provide a Typer CLI runner."""
    from typer.testing import CliRunner
    return CliRunner()


# ============================================================================
# Embedding Fixtures
# ============================================================================


@pytest.fixture
def sample_embedding() -> list[float]:
    """Provide a sample embedding vector."""
    return [0.1] * 3072  # text-embedding-3-large dimensions


@pytest.fixture
def mock_openai_embeddings():
    """Mock OpenAI embeddings API."""
    with patch("openai.OpenAI") as mock_client:
        mock_embedding = Mock()
        mock_embedding.data = [Mock(embedding=[0.1] * 3072)]
        
        mock_client.return_value.embeddings.create.return_value = mock_embedding
        yield mock_client
